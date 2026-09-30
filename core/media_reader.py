"""
多格式文件解析适配器: 把上传的文件统一打包成模型可读的 Content Block 列表。

支持范围:
- 图片: jpg/jpeg/png/webp/bmp, 以及 heic/heif (需 pillow-heif)
- PDF : 按页数采样关键页并渲染为 PNG
- 纯文本类: txt/md/json/log —— 直接作为文本送入 (带字符上限)
- Office: docx (含表格) —— 提取文本后送入

超出支持范围时明确报错，而不是静默产出垃圾数据。
"""
import base64
import hashlib
import io
import mimetypes
import shutil
from pathlib import Path
from typing import Dict, Any, List, Tuple

from .config import Config
from .pdf_processor import PDFProcessor

SUPPORTED_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
SUPPORTED_HEIF_EXTS = {".heic", ".heif"}
SUPPORTED_PDF_EXTS = {".pdf"}
SUPPORTED_TEXT_EXTS = {".txt", ".md", ".json", ".log"}
SUPPORTED_OFFICE_EXTS = {".docx"}

IMAGE_MIME = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
}


class MediaReader:
    @staticmethod
    def calculate_file_hash(file_bytes: bytes) -> str:
        return hashlib.sha256(file_bytes).hexdigest()[:16]

    @classmethod
    def save_and_archive(cls, source_path: Path) -> Tuple[str, str]:
        """将上传的文件按哈希值轻量归档到本地 attachments 目录 (内容相同则天然去重)

        目录内的文件**就地改名、不再复制**：机器人下载的原始文件本来就先落在
        attachments/ 下（见 wechat_bot.js），再复制一份的话每张单据在磁盘上占两个
        副本、长期只增不减。目录外的（样例、测试用的临时文件）仍然复制，不动调用方
        的东西——那些文件不属于我们。
        """
        if not source_path.exists():
            raise FileNotFoundError(f"未找到源文件: {source_path}")

        file_bytes = source_path.read_bytes()
        file_hash = cls.calculate_file_hash(file_bytes)
        ext = source_path.suffix.lower()

        target_name = f"{file_hash}{ext}"
        target_path = Config.ATTACHMENTS_DIR / target_name

        in_place = source_path.parent.resolve() == Config.ATTACHMENTS_DIR.resolve()
        if not target_path.exists():
            if in_place:
                shutil.move(str(source_path), str(target_path))
            else:
                shutil.copy2(source_path, target_path)
        elif in_place and source_path.resolve() != target_path.resolve():
            # 同样的内容已经有一份了，把这份多出来的删掉，别再留第二个副本
            source_path.unlink(missing_ok=True)

        return f"attachments/{target_name}", file_hash

    @classmethod
    def build_model_payloads(cls, file_path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
        """
        把文件打包为模型可读的 Content Block 列表。

        :return: (content_blocks, meta)
                 meta = {"kind": "image"|"pdf"|"text", "total_pages": int|None,
                         "sampled_pages": list|None, "chars": int|None}
        """
        ext = file_path.suffix.lower()

        if ext in SUPPORTED_IMAGE_EXTS or ext in SUPPORTED_HEIF_EXTS:
            raw, mime = cls._read_image(file_path, ext)
            return [cls._image_block(raw, mime)], {"kind": "image", "total_pages": None,
                                                   "sampled_pages": None, "chars": None}

        if ext in SUPPORTED_PDF_EXTS:
            total_pages, pages, images = PDFProcessor.render_sampled_pages(file_path)
            blocks = [cls._image_block(png, "image/png") for png in images]
            return blocks, {"kind": "pdf", "total_pages": total_pages,
                            "sampled_pages": pages, "chars": None}

        text = cls._read_text_like(file_path, ext)
        if text is not None:
            block = {"type": "text", "text": f"[文件内容: {file_path.name}]\n{text}"}
            return [block], {"kind": "text", "total_pages": None,
                             "sampled_pages": None, "chars": len(text)}

        supported = sorted(SUPPORTED_IMAGE_EXTS | SUPPORTED_HEIF_EXTS | SUPPORTED_PDF_EXTS
                           | SUPPORTED_TEXT_EXTS | SUPPORTED_OFFICE_EXTS)
        raise ValueError(f"不支持的文件格式: {ext or '(无扩展名)'}。本助手支持: {', '.join(supported)}")

    # ------------------------------------------------------------ 各格式读取
    @staticmethod
    def _read_image(file_path: Path, ext: str) -> Tuple[bytes, str]:
        if ext in SUPPORTED_HEIF_EXTS:
            try:
                import pillow_heif
                from PIL import Image
                pillow_heif.register_heif_opener()
            except ImportError as e:
                raise RuntimeError("读取 HEIC/HEIF 需要 pillow-heif，请执行: pip install pillow-heif") from e
            buf = io.BytesIO()
            with Image.open(file_path) as img:
                img.convert("RGB").save(buf, format="PNG")
            return buf.getvalue(), "image/png"

        mime = IMAGE_MIME.get(ext) or mimetypes.guess_type(str(file_path))[0] or "image/jpeg"
        return file_path.read_bytes(), mime

    @staticmethod
    def _read_text_like(file_path: Path, ext: str) -> "str | None":
        if ext in SUPPORTED_TEXT_EXTS:
            return file_path.read_text(encoding="utf-8", errors="replace")[:Config.MAX_TEXT_CHARS]

        if ext == ".docx":
            try:
                from docx import Document
            except ImportError as e:
                raise RuntimeError("读取 .docx 需要 python-docx，请执行: pip install python-docx") from e
            doc = Document(str(file_path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            return "\n".join(parts)[:Config.MAX_TEXT_CHARS]

        return None

    @staticmethod
    def _image_block(raw: bytes, mime: str) -> Dict[str, Any]:
        b64_str = base64.b64encode(raw).decode("utf-8")
        return {
            "type": "image_url",
            "image_url": {
                "url": f"data:{mime};base64,{b64_str}",
                "detail": "high",  # 保证发票等小字号高精度 OCR
            }
        }
