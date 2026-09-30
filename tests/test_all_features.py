"""
微信智能文档与账目助手 - 核心自动化测试集 (Pytest)
覆盖: 混合数据库与 JSON1 查询、真实 PDF 取页与渲染、上下文窗口口径、
      只读 SQL 安全防御、多格式媒体打包、清空历史前的备份完整性
"""
import json
import tempfile
from pathlib import Path

import pytest

from core.db import DatabaseManager
from core.pdf_processor import PDFProcessor
from core.context_compressor import ContextCompressor
from core.native_tools import NativeToolRegistry
from core.config import Config
from core.agent_runner import AgentRunner
from core.media_reader import MediaReader


@pytest.fixture
def temp_db():
    """提供隔离的临时数据库环境"""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield DatabaseManager(Path(tmpdir) / "test.db")


def make_pdf(path: Path, pages: int) -> Path:
    """用 PyMuPDF 生成一份真实的多页 PDF"""
    import fitz

    doc = fitz.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"Contract page {i + 1}")
    doc.save(str(path))
    doc.close()
    return path


# 1. 混合存储与 SQLite JSON1 穿透查询
def test_db_insert_and_json1_query(temp_db):
    id1 = temp_db.insert_document(
        doc_type="invoice", title="差旅车费发票", main_amount=350.0,
        main_date="2026-08-15", main_entity="中铁客运",
        extra_data={"tax_rate": 0.09, "tax_amount": 28.87, "seat": "二等座"},
    )
    id2 = temp_db.insert_document(
        doc_type="contract", title="材料采购年度框架合同", main_amount=120000.0,
        main_date="2026-08-01", main_entity="中铁二局物资部",
        extra_data={"valid_years": 2, "settlement": "月结"},
    )
    assert id1 > 0 and id2 > 0

    assert temp_db.query_documents("SELECT SUM(main_amount) AS total FROM documents;")[0]["total"] == 120350.0

    res_json = temp_db.query_documents(
        "SELECT * FROM documents WHERE json_extract(extra_json, '$.tax_rate') = 0.09;"
    )
    assert len(res_json) == 1
    assert res_json[0]["main_entity"] == "中铁客运"
    assert res_json[0]["extra_json"]["seat"] == "二等座"


# 3. 多页 PDF 取页 (只读开头几页，真实读取页数 + 真实渲染)
def test_pdf_page_sampling_real(tmp_path):
    # 短文档 (3 页) 不超过上限，全读
    assert PDFProcessor.sample_pdf_pages(3) == [1, 2, 3]
    # 长合同 (30 页) 只读前 3 页，后面的不读
    assert PDFProcessor.sample_pdf_pages(30) == [1, 2, 3]

    six_pages = make_pdf(tmp_path / "contract_6p.pdf", 6)
    assert PDFProcessor.get_page_count(six_pages) == 6

    total, pages, images = PDFProcessor.render_sampled_pages(six_pages, dpi=72)
    assert total == 6
    assert pages == [1, 2, 3]
    assert len(images) == 3
    assert all(img.startswith(b"\x89PNG") for img in images)   # 确实是渲染出的 PNG

    # 无法解析的文件必须明确报错，而不是静默产出垃圾数据
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4\nnot a real pdf")
    with pytest.raises(ValueError):
        PDFProcessor.get_page_count(broken)


# 4. 自动上下文批量整理与物理留底保障机制
def test_context_auto_compression_and_no_amnesia(temp_db):
    compressor = ContextCompressor(temp_db, max_raw_rounds=2, compression_threshold_rounds=3)

    compressor.save_message("user", "你好，我想登记一张发票")
    compressor.save_message("assistant", "请发送发票图片")
    compressor.save_message("user", "发票金额是 5800 元，开票方是顺丰速运")
    compressor.save_message("assistant", "已确认顺丰速运 5800 元发票")
    compressor.save_message("user", "帮我查一下上月总共花了多少")
    compressor.save_message("assistant", "正在为您查询...")
    compressor.save_message("user", "还有中铁二局的合同在哪？")
    compressor.save_message("assistant", "中铁二局合同编号为 HT-2026-01")

    # 未压缩 4 轮 > 阈值 3 轮，应自动触发批量整理
    context_msgs = compressor.build_model_context()

    system_content = context_msgs[0]["content"]
    # 坚决不做不准确的正则弱事实抽取，避免提示词污染
    assert "长期记忆事实锚点" not in system_content
    # 坚决不做多余的 LLM 摘要：单据在数据库随时可查，不占 System 前缀
    assert "阶段要点" not in system_content

    # 窗口按轮计: 保留最近 2 轮 = 4 条对话消息，且首条必须是 user（不能切出孤儿回复）
    raw_msgs = [m for m in context_msgs if m["role"] != "system"]
    assert len(raw_msgs) == 4
    assert raw_msgs[0]["role"] == "user"
    assert "HT-2026-01" in raw_msgs[-1]["content"]

    # 较早的消息从提示词中移出，节省 Token 并净化上下文
    whole_context = json.dumps(context_msgs, ensure_ascii=False)
    assert "请发送发票图片" not in whole_context

    # 关键防失忆底线：chat_messages 物理全量持久化留底，历史对话永久可查，绝不真删
    with temp_db.get_connection() as conn:
        all_stored = conn.execute("SELECT content FROM chat_messages ORDER BY id ASC;").fetchall()
    assert len(all_stored) == 8
    assert any("请发送发票图片" in r["content"] for r in all_stored)


# 5. 原生 Tool Calling 与只读 SQL 安全防御
def test_native_tool_calling_and_sql_safety(temp_db):
    tools = NativeToolRegistry(temp_db)
    temp_db.insert_document("invoice", "差旅费", 200.0, "2026-08-01", "滴滴出行")

    res = tools.execute_tool("query_documents_by_sql", {
        "sql": "SELECT main_amount, main_entity FROM documents WHERE main_entity='滴滴出行';"
    })
    assert res["status"] == "success"
    assert res["data"][0]["main_amount"] == 200.0

    # 破坏性与越权语句必须全部被拒绝（只读连接 + authorizer 白名单）
    for danger_sql in [
        "DROP TABLE documents;",
        "DELETE FROM documents;",
        "UPDATE documents SET main_amount = 0;",
        "ATTACH DATABASE '/tmp/evil.db' AS evil;",
    ]:
        danger = tools.execute_tool("query_documents_by_sql", {"sql": danger_sql})
        assert danger["status"] == "error"
        assert "安全拦截" in danger["message"]

    summary_res = tools.execute_tool("get_spending_summary", {"year_month": "2026-08"})
    assert summary_res["status"] == "success"
    assert summary_res["total_amount"] == 200.0


# 6. 上下文前缀稳定性 (常规轮次严格单调追加，前缀逐字节不变)
def test_context_prefix_stability(temp_db):
    compressor = ContextCompressor(temp_db, max_raw_rounds=3, compression_threshold_rounds=8)

    hashes = []
    for i in range(3):
        compressor.save_message("user", f"第 {i + 1} 句话")
        compressor.save_message("assistant", f"回复 {i + 1}")
        hashes.append(compressor.calculate_prefix_hash(check_length=2))

    assert hashes[0] == hashes[1] == hashes[2]


# 7. 多格式文件打包 (图片直传 / PDF 采样渲染)
def test_multi_format_media_reader(tmp_path):
    png_file = tmp_path / "test_invoice.png"
    png_file.write_bytes(b"\x89PNG\r\n\x1a\nfake_png_data")

    blocks, meta = MediaReader.build_model_payloads(png_file)
    assert meta["kind"] == "image"
    assert len(blocks) == 1
    assert "data:image/png;base64," in blocks[0]["image_url"]["url"]
    assert blocks[0]["image_url"]["detail"] == "high"

    # 多页 PDF 只渲染前几页，后面的不进载荷
    pdf_file = make_pdf(tmp_path / "contract_8p.pdf", 8)
    pdf_blocks, pdf_meta = MediaReader.build_model_payloads(pdf_file)
    assert pdf_meta == {"kind": "pdf", "total_pages": 8, "sampled_pages": [1, 2, 3], "chars": None}
    assert len(pdf_blocks) == 3
    assert all("data:image/png;base64," in b["image_url"]["url"] for b in pdf_blocks)

    # 哈希归档
    rel_path, f_hash = MediaReader.save_and_archive(png_file)
    assert len(f_hash) == 16
    assert rel_path.startswith("attachments/")


# 8. 单据上传 -> 多模态结构化录入端到端
def test_handle_document_upload_integration(temp_db, archive_dir):
    runner = AgentRunner(temp_db)
    test_img = archive_dir / "shopping_receipt.png"
    test_img.write_bytes(b"\x89PNG\r\n\x1a\nsimulated_receipt_image_bytes")

    mock_vision_data = {
        "doc_type": "order",
        "title": "淘宝订单-机械键盘",
        "main_amount": 399.0,
        "main_date": "2026-09-26",
        "main_entity": "罗技旗舰店",
        "amount_without_tax": 0.0,
        "tax_amount": 0.0,
        "summary": "罗技机械键盘购物实付 399 元",
    }

    reply = runner.handle_document_upload(
        file_path=str(test_img),
        user_text="帮我把这个键盘记一下账",
        mock_llm_response=mock_vision_data,
    )

    assert "已归档" in reply
    assert "¥399.00" in reply
    assert "罗技旗舰店" in reply

    docs = temp_db.query_documents("SELECT * FROM documents WHERE main_entity='罗技旗舰店'")
    assert len(docs) == 1
    assert docs[0]["main_amount"] == 399.0
    assert docs[0]["title"] == "淘宝订单-机械键盘"


def test_handle_document_upload_works_when_file_starts_inside_archive_dir(temp_db, tmp_path, monkeypatch):
    """回归测试：文件落在归档目录内时，归档是 move 而非 copy。

    微信发来的图/文件就是这种落盘方式（bot 下载进 data/attachments/ 再 POST 该路径）。
    修这个 bug 之前，save_and_archive 把原文件移走后，调用方仍拿原路径去读，
    FileNotFoundError 逃过 ValueError/RuntimeError 处理器，微信上传全线 500。
    原来的测试只测 save_and_archive 本身，而集成测试用的文件在归档目录外（走 copy 分支），
    永远碰不到这条路径。
    """
    from core.config import Config
    archive = tmp_path / "attachments"
    archive.mkdir()
    monkeypatch.setattr(Config, "ATTACHMENTS_DIR", archive)

    runner = AgentRunner(temp_db)
    incoming = archive / "wx_img_1730000000.png"   # 模拟 bot 下载到归档目录里的微信图片
    incoming.write_bytes(b"\x89PNG\r\n\x1a\nsimulated_receipt_image_bytes")

    reply = runner.handle_document_upload(
        file_path=str(incoming),
        user_text="帮我记一下",
        mock_llm_response={
            "doc_type": "order", "title": "淘宝订单-机械键盘",
            "main_amount": 399.0, "main_date": "2026-09-26",
            "main_entity": "罗技旗舰店", "amount_without_tax": 0.0,
            "tax_amount": 0.0, "summary": "罗技机械键盘购物实付 399 元",
        },
    )

    assert "已归档" in reply
    assert "¥399.00" in reply
    # 原文件被移走是预期行为，但归档后的那份必须存在且可读
    assert not incoming.exists()
    assert len(list(archive.glob("*.png"))) == 1
    docs = temp_db.query_documents("SELECT * FROM documents WHERE main_entity='罗技旗舰店'")
    assert len(docs) == 1


# 9. 工具调用闭环: 模型请求工具 -> 执行 -> 结果回灌
def test_tool_calling_loop_with_mocked_model(temp_db):
    runner = AgentRunner(temp_db)
    temp_db.insert_document("invoice", "办公用品", 88.0, "2026-08-09", "晨光文具")

    calls = []

    def fake_completion(messages, tools=None, temperature=0.1):
        calls.append(tools)
        if len(calls) == 1:
            return {
                "status": "success",
                "content": "",
                "tool_calls": [{
                    "id": "call_1",
                    "function": {
                        "name": "query_documents_by_sql",
                        "arguments": '{"sql": "SELECT SUM(main_amount) AS total FROM documents;"}',
                    },
                }],
            }
        return {"status": "success", "content": "本月累计支出 88.00 元。", "tool_calls": []}

    runner.llm.chat_completion = fake_completion
    reply = runner.handle_user_message("这个月花了多少？")

    assert reply == "本月累计支出 88.00 元。"
    assert len(calls) == 2          # 第一轮发起工具调用，第二轮给出结论
    assert calls[0] is not None     # 首轮带 tools schema


# 10. 单据检索工具: 关键字 / 类型 / 月份过滤
def test_find_documents_filters(temp_db):
    tools = NativeToolRegistry(temp_db)
    temp_db.insert_document("invoice", "办公用品发票", 1130.0, "2026-08-18", "上海晨光办公用品有限公司")
    temp_db.insert_document("ticket", "高铁车票", 553.0, "2026-08-05", "中国铁路")
    temp_db.insert_document("receipt", "餐饮小票", 88.5, "2026-07-20", "某某餐饮有限公司")

    by_keyword = tools.execute_tool("find_documents", {"keyword": "晨光"})
    assert by_keyword["count"] == 1
    assert by_keyword["documents"][0]["title"] == "办公用品发票"
    assert by_keyword["documents"][0]["has_file"] is False   # 未归档原件时如实告知

    assert tools.execute_tool("find_documents", {"doc_type": "ticket"})["count"] == 1
    assert tools.execute_tool("find_documents", {"year_month": "2026-08"})["count"] == 2
    assert tools.execute_tool("find_documents", {})["count"] == 3


# 11. 演示数据脚本: 灌入后可检索，且每条都登记了真实存在的原件
def test_seed_demo_data(temp_db, tmp_path, monkeypatch):
    from core.config import Config
    import seed_demo_data

    archive_dir = tmp_path / "attachments"
    archive_dir.mkdir()
    monkeypatch.setattr(Config, "ATTACHMENTS_DIR", archive_dir)
    monkeypatch.setattr(Config, "DATA_DIR", tmp_path)

    inserted = seed_demo_data.seed(temp_db)
    assert inserted >= 5

    docs = temp_db.query_documents("SELECT id, doc_type, file_path FROM documents;")
    assert len(docs) == inserted
    assert len({d["doc_type"] for d in docs}) >= 5          # 覆盖多种单据类型
    for d in docs:
        archived = tmp_path / d["file_path"]
        assert archived.is_file()
        assert archived.resolve().is_relative_to(archive_dir.resolve())

    # 已有数据时默认不重复灌入
    assert seed_demo_data.seed(temp_db) == 0


# 12. 单据类型归一化: 模型返回自由文本类型时也要落到统一分类，保证检索口径一致
def test_normalize_doc_type():
    from core.agent_runner import normalize_doc_type

    # 模型直接给对枚举值时原样保留
    assert normalize_doc_type("invoice") == "invoice"
    assert normalize_doc_type("ticket") == "ticket"

    # 实测中模型曾把发票返回成"增值税电子普通发票"，必须归一为 invoice
    assert normalize_doc_type("增值税电子普通发票", "增值税电子普通发票") == "invoice"
    assert normalize_doc_type("机票行程单", "航空运输电子客票行程单") == "ticket"
    assert normalize_doc_type("对账单", "8 月银行对账单") == "statement"
    assert normalize_doc_type("差旅费用报销单", "差旅报销") == "reimbursement"
    assert normalize_doc_type("材料送货单", "螺纹钢送货单") == "delivery"
    assert normalize_doc_type("采购合同", "机械设备采购合同") == "contract"
    assert normalize_doc_type("餐饮小票", "某某餐饮消费小票") == "receipt"
    assert normalize_doc_type("订单截图", "淘宝订单详情") == "order"

    # 无法判断时落到 other
    assert normalize_doc_type("something-weird", "无法识别的材料") == "other"


# 14. 文本类与 Office 文件的解析 (txt/docx)
def test_text_and_office_payloads(tmp_path):
    txt_file = tmp_path / "note.txt"
    txt_file.write_text("发票号码 02514789，价税合计 1130.00 元", encoding="utf-8")
    txt_blocks, txt_meta = MediaReader.build_model_payloads(txt_file)
    assert txt_meta["kind"] == "text"
    assert "1130.00" in txt_blocks[0]["text"]

    from docx import Document
    doc = Document()
    doc.add_paragraph("机械设备采购合同")
    doc.add_paragraph("甲方: 某某建设集团有限公司")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "合同总金额"
    table.rows[0].cells[1].text = "120000.00"
    docx_path = tmp_path / "contract.docx"
    doc.save(docx_path)
    docx_blocks, docx_meta = MediaReader.build_model_payloads(docx_path)
    assert docx_meta["kind"] == "text"
    assert "机械设备采购合同" in docx_blocks[0]["text"]
    assert "120000.00" in docx_blocks[0]["text"]      # 表格内容也要被提取

    # 表格类文件不在支持范围内，必须给出明确提示而不是静默处理
    for table_name in ("statement.csv", "statement.xlsx", "statement.xls"):
        table_file = tmp_path / table_name
        table_file.write_text("日期,对方,金额\n2026-08-01,某某建材,120000.00", encoding="utf-8")
        with pytest.raises(ValueError) as err:
            MediaReader.build_model_payloads(table_file)
        assert "不支持的文件格式" in str(err.value)
        assert ".pdf" in str(err.value)

    # 不支持的格式必须明确报错，而不是静默产出垃圾数据
    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    with pytest.raises(ValueError):
        MediaReader.build_model_payloads(legacy)

    # 4. 超长文档字数上限截断保护 (防止大部头超长文档撑爆模型上下文)
    huge_txt = tmp_path / "huge.txt"
    huge_txt.write_text("长文本测试" * 6000, encoding="utf-8")  # 30,000 字符
    huge_blocks, huge_meta = MediaReader.build_model_payloads(huge_txt)
    assert huge_meta["chars"] <= Config.MAX_TEXT_CHARS
    assert len(huge_blocks[0]["text"]) <= Config.MAX_TEXT_CHARS + 50


# 15. 工具结果行数上限 (防止大批量查询把上下文撑爆)
def test_tool_result_row_cap(temp_db, monkeypatch):
    from core.config import Config

    monkeypatch.setattr(Config, "MAX_TOOL_ROWS_IN_CONTEXT", 3)
    tools = NativeToolRegistry(temp_db)
    for i in range(6):
        temp_db.insert_document("invoice", f"单据{i}", 10.0, "2026-08-01", "某公司")

    capped = tools.execute_tool("query_documents_by_sql", {"sql": "SELECT * FROM documents;"})
    assert capped["count"] == 6            # 真实总行数如实回报
    assert len(capped["data"]) == 3        # 只回灌前 3 行
    assert capped["truncated"] is True
    assert "LIMIT" in capped["note"]

    # 未超限时不做截断
    intact = tools.execute_tool("query_documents_by_sql", {"sql": "SELECT * FROM documents LIMIT 2;"})
    assert "truncated" not in intact
    assert len(intact["data"]) == 2


# 16. HEIC 图片 (iPhone 照片) 读取: 转码为 PNG 送模型，原件按原扩展名归档
def test_heic_image_handling(tmp_path):
    pillow_heif = pytest.importorskip("pillow_heif", reason="需要 pillow-heif 才能构造 HEIC 用例")
    from PIL import Image

    pillow_heif.register_heif_opener()
    source = tmp_path / "src.png"
    Image.new("RGB", (40, 30), (255, 0, 0)).save(source)

    heic_path = tmp_path / "iphone_photo.heic"
    with Image.open(source) as img:
        img.convert("RGB").save(heic_path, format="HEIF")

    blocks, meta = MediaReader.build_model_payloads(heic_path)
    assert meta["kind"] == "image"
    assert blocks[0]["image_url"]["url"].startswith("data:image/png;base64,")   # 已转码

    rel_path, _ = MediaReader.save_and_archive(heic_path)
    assert rel_path.endswith(".heic")      # 原件保留原始格式


# 17. 疑似重复凭证检测: 只提示、不自动合并; 依据不足时宁可返回空
def test_similar_document_detection(temp_db):
    first = temp_db.insert_document("invoice", "餐饮发票", 88.0, "2026-08-15", "某某餐饮有限公司")

    # 同主体(简称)、同金额、日期相差 1 天 -> 判定为疑似同一笔的不同凭证
    hits = temp_db.find_similar_documents("某某餐饮公司", 88.0, "2026-08-16")
    assert [h["id"] for h in hits] == [first]

    # 主体不同 -> 不猜
    assert temp_db.find_similar_documents("另一家科技有限公司", 88.0, "2026-08-16") == []
    # 金额差异超出容差 (团购 88 vs 发票 500) -> 不猜
    assert temp_db.find_similar_documents("某某餐饮有限公司", 500.0, "2026-08-16") == []
    # 日期超出窗口 -> 不猜
    assert temp_db.find_similar_documents("某某餐饮有限公司", 88.0, "2026-12-30") == []
    # 主体未知 (识别失败) -> 不猜
    assert temp_db.find_similar_documents("未知主体", 88.0, "2026-08-15") == []


# 18. 重复提示只提示不合并: 回复里要有提示，且库里仍是两条独立记录
def test_upload_reports_possible_duplicate_without_merging(temp_db, archive_dir, monkeypatch):
    from core.config import Config
    monkeypatch.setattr(Config, "DATA_DIR", archive_dir.parent)

    runner = AgentRunner(temp_db)
    # 发现问题后是把情况回给模型、由模型组织措辞，所以这里把"模型那一句"打桩：
    # 既避免测试真去调 API，也把断言钉死在**程序有没有发现重复**这件事上。
    seen_warnings = []

    def fake_ask(just_said, warnings):
        seen_warnings.append(list(warnings))
        return "这好像跟前面那张是同一笔，我没有自动合并，你确认下。"

    monkeypatch.setattr(runner, "_ask_model_about_warnings", fake_ask)

    doc_data = {
        "doc_type": "invoice", "title": "餐饮发票", "main_amount": 88.0,
        "main_date": "2026-08-15", "main_entity": "某某餐饮有限公司",
        "reply": "餐饮发票 ¥88.00",
    }

    first_img = archive_dir / "receipt_a.png"
    first_img.write_bytes(b"\x89PNG\r\n\x1a\nfirst")
    runner.handle_document_upload(str(first_img), mock_llm_response=doc_data)
    assert seen_warnings == []                          # 第一条没有任何可疑之处

    second_img = archive_dir / "receipt_b.png"
    second_img.write_bytes(b"\x89PNG\r\n\x1a\nsecond")
    second_reply = runner.handle_document_upload(str(second_img), mock_llm_response=dict(doc_data))
    assert any("很接近" in w for w in seen_warnings[-1])  # 第二条把重复报给了模型
    assert "没有自动合并" in second_reply

    # 关键: 只是提示，没有真的合并 —— 仍是两条独立记录
    assert len(temp_db.query_documents("SELECT id FROM documents;")) == 2


def test_missing_date_is_reported_not_faked(temp_db, archive_dir, monkeypatch):
    """模型没读到日期时不能拿今天顶上——那是编数据。要把这件事回给模型去问用户。"""
    from core.config import Config
    monkeypatch.setattr(Config, "DATA_DIR", archive_dir.parent)

    runner = AgentRunner(temp_db)
    seen = []

    def fake_ask(just_said, warnings):
        seen.append(list(warnings))
        return "这张凭证上没看到日期，是哪天的呀？"

    monkeypatch.setattr(runner, "_ask_model_about_warnings", fake_ask)

    img = archive_dir / "nodate.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\nnodate")
    reply = runner.handle_document_upload(str(img), mock_llm_response={
        "doc_type": "order", "title": "微信支付凭证", "main_amount": 88.0,
        "main_date": "", "main_entity": "某某超市", "reply": "微信支付凭证 ¥88.00",
    })

    assert any("日期" in w for w in seen[-1])
    assert temp_db.query_documents("SELECT main_date FROM documents;")[0]["main_date"] == ""
    assert "哪天" in reply

    # 编号要留在内部记录里，模型才知道用户事后补充时该改哪一条；
    # 但它不能出现在发给用户的回执里
    sent = temp_db.query_documents("SELECT content FROM chat_messages WHERE role='user'")
    assert any("已归档为 #" in m["content"] for m in sent)
    assert "#" not in reply


# 19. 批量整理游标正确推进且前缀哈希保持稳定
def test_context_compaction_cursor_and_prefix_stability(temp_db):
    compressor = ContextCompressor(temp_db, max_raw_rounds=2, compression_threshold_rounds=3)

    # 1. 保存前 4 条消息（= 2 轮），验证前缀哈希逐字节稳定
    for i in range(4):
        compressor.save_message("user" if i % 2 == 0 else "assistant", f"消息 {i}")
    hash1 = compressor.calculate_prefix_hash(check_length=1)

    # 2. 追加到 6 条（= 3 轮，未超过阈值 3 轮），System Prompt 前缀逐字节绝对不变
    compressor.save_message("user", "消息 4")
    compressor.save_message("assistant", "消息 5")
    hash2 = compressor.calculate_prefix_hash(check_length=1)
    assert hash1 == hash2

    # 3. 达到 8 条（= 4 轮 > 阈值 3 轮），触发整理，游标推进
    compressor.save_message("user", "消息 6")
    compressor.save_message("assistant", "消息 7")
    context = compressor.build_model_context()

    state = compressor.get_context_state()
    # 保留最近 2 轮（第 5~8 条），游标停在前面那条回复上 = 第 4 条消息的 ID
    assert state["compressed_up_to_id"] == 4
    # 系统提示词前缀依然保持纯净，哈希未变
    hash3 = compressor.calculate_prefix_hash(check_length=1)
    assert hash1 == hash3


# 20. SQL 工具不得把内部归档路径暴露给模型
# (实测踩到过: 模型从 SQL 结果里读到 file_path 后，把服务器路径念给用户，而那对用户毫无意义)
def test_sql_tool_hides_internal_file_path(temp_db, tmp_path, monkeypatch):
    from core.config import Config
    monkeypatch.setattr(Config, "DATA_DIR", tmp_path)

    image = tmp_path / "receipt.png"
    image.write_bytes(b"\x89PNG\r\n\x1a\nreceipt")
    rel_path, _ = MediaReader.save_and_archive(image)
    temp_db.insert_document("invoice", "办公用品发票", 1130.0, "2026-08-18", "某公司",
                            file_path=rel_path)

    tools = NativeToolRegistry(temp_db)
    result = tools.execute_tool("query_documents_by_sql", {"sql": "SELECT * FROM documents;"})
    row = result["data"][0]
    assert "file_path" not in row        # 内部路径不暴露给模型
    assert row["has_file"] is True       # 只暴露"有没有原件"


# 21. System prompt 自身也不得携带归档路径
# (实测踩到过: 上面那条测试守住了工具层，但提示词里写着"说明原件归档在本机 data/attachments
#  目录下"，等于把刚剥掉的路径又加回去，直接指示模型做与 README 承诺相反的事)
def test_system_prompt_does_not_leak_archive_path():
    prompt = ContextCompressor.STATIC_BASE_SYSTEM
    # 只查真正的路径形态；"归档路径"这个词本身可以出现——提示词正是用它来禁止念路径的
    for token in ("data/attachments", "attachments/", "/Users/", "file_path", "~/", "/home/"):
        assert token not in prompt, f"system prompt 里不该出现「{token}」"
    assert "只能发文字" in prompt        # 通道能力限制仍要如实说明
    # 编号只做内部把手，不出现给用户看的内容里——和"让用户记编号"是相反的
    assert "一个字都不要提" in prompt


# 22. 工具帧不得挤占记忆窗口，也不得切出"没有对应提问的回复"
# (实测踩到过: 窗口与阈值按"消息条数"计，而每次工具调用都会落一条 tool 帧，
#  于是工具帧既占阈值又占窗口——真实用法下 20 轮的窗口被压到 4~14 轮，
#  且 276 次整理里上下文以一条孤儿 assistant 回复开头。)
def test_tool_frames_do_not_shrink_memory_window(temp_db):
    compressor = ContextCompressor(temp_db, max_raw_rounds=3, compression_threshold_rounds=4)

    # 6 轮，每轮都带 3 条工具结果帧（真实用法: 上传单据与查账都会调工具）
    for i in range(6):
        compressor.save_message("user", f"第 {i} 问")
        for _ in range(3):
            compressor.save_message("tool", '{"status":"success","data":[]}')
        compressor.save_message("assistant", f"第 {i} 答")

    context = compressor.build_model_context()
    dialogue = [m for m in context if m["role"] != "system"]

    # 窗口是 3 轮 = 6 条对话消息，不能被 18 条工具帧挤掉
    assert len(dialogue) == 6
    # 首条必须是 user: 不能出现"留下了回复、却砍掉它对应的提问"的孤儿消息
    assert dialogue[0]["role"] == "user"
    assert dialogue[0]["content"] == "第 3 问"
    assert dialogue[-1]["content"] == "第 5 答"
    assert compressor.last_context_rounds == 3

    # 工具帧本身仍物理留底（可回溯模型当时查了什么），只是不占窗口、不影响阈值
    with temp_db.get_connection() as conn:
        tool_rows = conn.execute(
            "SELECT COUNT(*) FROM chat_messages WHERE role='tool'").fetchone()[0]
    assert tool_rows == 18


# 23. 清空历史前的备份必须一条不漏
# (实测踩到过: backup() 走了给模型用的查询通道，被 MAX_QUERY_ROWS 截断——
#  250 条消息只导出 200 条，随后 DELETE 一执行，后 50 条永久丢失且无备份)
def test_chat_backup_is_not_truncated_by_query_row_cap(temp_db, tmp_path, monkeypatch):
    from core.config import Config
    import reset_chat

    monkeypatch.setattr(Config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(Config, "MAX_QUERY_ROWS", 5)      # 把上限压小，等价于线上的 200 行

    with temp_db.get_connection() as conn:
        for i in range(12):
            conn.execute("INSERT INTO chat_messages (role, content) VALUES ('user', ?)",
                         (f"第 {i} 条",))
        conn.commit()

    path = reset_chat.backup(temp_db)
    written = json.loads(path.read_text(encoding="utf-8"))

    assert len(written["messages"]) == 12                 # 全部导出，不受 MAX_QUERY_ROWS 影响
    assert written["messages"][0]["content"] == "第 0 条"
    assert written["messages"][-1]["content"] == "第 11 条"


# 24. 归档不再留两个副本
# (机器人把下载的原始文件先落在 attachments/ 下，归档时又复制一份哈希副本——
#  每张单据在磁盘上占两份，长期只增不减)
def test_archive_moves_incoming_file_instead_of_copying(tmp_path, monkeypatch):
    from core.config import Config
    archive = tmp_path / "attachments"
    archive.mkdir()
    monkeypatch.setattr(Config, "ATTACHMENTS_DIR", archive)

    # 1) 落在归档目录内的（机器人下载的原始文件）-> 就地改名，只留一份
    incoming = archive / "wx_img_1730000000.jpg"
    incoming.write_bytes(b"\x89PNG\r\n\x1a\nreceipt-bytes")
    rel, digest = MediaReader.save_and_archive(incoming)
    assert not incoming.exists()
    assert (archive / f"{digest}.jpg").is_file()
    assert rel == f"attachments/{digest}.jpg"

    # 2) 目录外的（样例文件）-> 复制，不动调用方的东西
    outside = tmp_path / "sample.png"
    outside.write_bytes(b"\x89PNG\r\n\x1a\nsample-bytes")
    _, digest2 = MediaReader.save_and_archive(outside)
    assert outside.is_file()
    assert (archive / f"{digest2}.png").is_file()

    # 3) 同一份内容再进来一次 -> 已有哈希副本，不该留下第二个
    again = archive / "wx_img_1730000001.jpg"
    again.write_bytes(b"\x89PNG\r\n\x1a\nreceipt-bytes")
    MediaReader.save_and_archive(again)
    assert not again.exists()
    assert len(list(archive.glob("*.jpg"))) == 1



def test_upload_rejects_path_outside_archive_dir(temp_db, archive_dir, tmp_path):
    """附件目录外的路径必须被拒。

    /api/message 没有鉴权，filePath 由客户端 POST 传入。不校验就会变成任意文件读取口：
    POST 一个 .env 或 /etc/passwd 的路径，模型会把内容原样回显在回复里。
    """
    outside = tmp_path / "secret.txt"
    outside.write_text("OPENAI_API_KEY=sk-should-never-be-echoed", encoding="utf-8")

    runner = AgentRunner(temp_db)
    reply = runner.handle_document_upload(str(outside), user_text="读一下这个")

    assert "没找到这条消息带过来的文件" in reply
    assert "sk-should-never-be-echoed" not in reply
    assert temp_db.query_documents("SELECT * FROM documents") == []
